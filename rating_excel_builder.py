"""
rating_excel_builder.py
------------------------------------------------------------------
Builds the "Rating Specification" Excel workbook (Sheet 1 = Rating
Steps, Sheet 2 = Rating Tables) in the EXACT same visual format as
the sample file the team uses (M_GG_PrivateMotor_RatingSpecification).

WHY THIS IS A SEPARATE FILE:
This project is shared between 2 contributors on GitHub. To avoid
merge conflicts, this file is 100% new/independent -- it does not
edit app.py, config.py, rag.py, retriver.py or ingest.py. It only
gets IMPORTED by the new export script (export_rating_to_excel.py).

HOW THE STYLING WAS DERIVED:
The colors below are not "made up" -- they were extracted directly
from the sample workbook's theme (xl/theme/theme1.xml):
    accent1 (#145F82) + 80% tint  -> #C1E5F5  (header row band)
    accent2 (#E87331) + 80% tint  -> #FAE3D6  (Step-title band)
Font = Calibri, thin black borders everywhere, same column widths
as the sample "Rating Steps" / "Rating Tables" sheets.
------------------------------------------------------------------
"""

from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


# ------------------------------------------------------------------
# STYLE CONSTANTS  (copied 1:1 from the sample workbook so every new
# product export looks identical to it -- do not change unless the
# sample template itself changes)
# ------------------------------------------------------------------
FONT_NAME = "Calibri"
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True)
COLUMN_HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True)
STEP_TITLE_FONT = Font(name=FONT_NAME, size=11, bold=True)
BODY_FONT = Font(name=FONT_NAME, size=11, bold=False)
OUTPUT_FONT = Font(name=FONT_NAME, size=11, bold=True)  # "Output" col (P1, L1..) is bold in sample

HEADER_FILL = PatternFill("solid", fgColor="C1E5F5")   # column header row band (light blue)
STEP_FILL = PatternFill("solid", fgColor="FAE3D6")      # "Step N - ..." band (light peach)
TABLE_TITLE_FILL = PatternFill("solid", fgColor="C1E5F5")  # table name band on Rating Tables sheet

THIN = Side(style="thin", color="FF000000")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Rating Steps sheet column widths (matches sample: B..F)
RATING_STEPS_COL_WIDTHS = {"A": 9.11, "B": 30.11, "C": 42.55, "D": 28.55, "E": 35.66, "F": 15.55}

# Column headers for the Rating Steps sheet (row 3 in the sample)
STEP_SHEET_HEADERS = ["Premium Component", "Condition", "Rate Value Reference", "Calculation Logic", "Output"]


# ------------------------------------------------------------------
# Data contract expected from the extractor (product_rating_extractor.py)
# ------------------------------------------------------------------
# steps = [
#   {
#     "step_title": "Step 1 - Base Premium Computation",
#     "rows": [
#         {
#           "component": "Comprehensive Cover Base Premium",   # "" = continuation of the row above (merges cell)
#           "condition": "If Type of Cover = Comprehensive ...",
#           "rate_reference": "Refer 'Base Amount' using ...", # "" = continuation
#           "calc_logic": "Base Premium = (11.5% * $30,000) + Base Amount",
#           "output": "P1"                                      # "" = continuation
#         },
#         ...
#     ]
#   },
#   ...
# ]
#
# tables = [
#   {
#     "table_name": "Base Amount Lookup",
#     "columns": ["Type of Cover", "Country", "Vehicle CC", "Base Amount"],
#     "rows": [["Comprehensive", "Trinidad", "Up to 2200cc", 1150], ...]
#   },
#   ...
# ]
# ------------------------------------------------------------------


def _merge_repeated_column(ws: Worksheet, col_letter: str, start_row: int, end_row: int) -> None:
    """
    Vertically merges consecutive cells in a column that share the same
    (non-blank) value written on the first row of the run. This mirrors
    how the sample sheet merges e.g. B5:B8 (one "Premium Component" label
    covering 4 condition rows) or F5:F12 (one "Output" code P1 covering
    both Comprehensive and TPFT sub-conditions).

    The extractor marks a "continuation" row by leaving the field blank,
    so we just scan for contiguous blank runs following a filled cell.
    """
    row = start_row
    while row <= end_row:
        value = ws[f"{col_letter}{row}"].value
        if value not in (None, ""):
            merge_end = row
            # Extend the merge while the *next* rows are blank continuations
            while merge_end + 1 <= end_row and ws[f"{col_letter}{merge_end + 1}"].value in (None, ""):
                merge_end += 1
            if merge_end > row:
                ws.merge_cells(f"{col_letter}{row}:{col_letter}{merge_end}")
                # Re-apply alignment/border to the merged block (openpyxl only
                # keeps the top-left cell's style after merge, but border needs
                # to be reapplied to look continuous)
                for r in range(row, merge_end + 1):
                    ws[f"{col_letter}{r}"].border = THIN_BORDER
            row = merge_end + 1
        else:
            row += 1


def _write_rating_steps_sheet(ws: Worksheet, product_name: str, steps: List[Dict[str, Any]]) -> None:
    """Writes Sheet 1 exactly like the sample 'Rating Steps' tab."""

    for col, width in RATING_STEPS_COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    # Title band (merged B1:F2 in the sample)
    ws.merge_cells("B1:F2")
    title_cell = ws["B1"]
    title_cell.value = f"{product_name} - Premium Computation Logic"
    title_cell.font = TITLE_FONT
    title_cell.alignment = WRAP_CENTER

    # Column header row (row 3 in the sample)
    header_row = 3
    for offset, header_text in enumerate(STEP_SHEET_HEADERS):
        col_letter = get_column_letter(2 + offset)  # B, C, D, E, F
        cell = ws[f"{col_letter}{header_row}"]
        cell.value = header_text
        cell.font = COLUMN_HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = THIN_BORDER

    current_row = header_row + 1

    for step in steps:
        # Step title band, merged B:F (e.g. "Step 1 - Base Premium Computation")
        step_row = current_row
        ws.merge_cells(f"B{step_row}:F{step_row}")
        step_cell = ws[f"B{step_row}"]
        step_cell.value = step.get("step_title", "")
        step_cell.font = STEP_TITLE_FONT
        step_cell.fill = STEP_FILL
        step_cell.alignment = WRAP_LEFT
        step_cell.border = THIN_BORDER
        current_row += 1

        rows = step.get("rows", [])
        block_start = current_row

        for row_data in rows:
            ws[f"B{current_row}"] = row_data.get("component", "")
            ws[f"C{current_row}"] = row_data.get("condition", "")
            ws[f"D{current_row}"] = row_data.get("rate_reference", "")
            ws[f"E{current_row}"] = row_data.get("calc_logic", "")
            ws[f"F{current_row}"] = row_data.get("output", "")

            for col_letter, is_output in (("B", False), ("C", False), ("D", False), ("E", False), ("F", True)):
                cell = ws[f"{col_letter}{current_row}"]
                cell.font = OUTPUT_FONT if is_output else BODY_FONT
                cell.alignment = WRAP_LEFT
                cell.border = THIN_BORDER

            current_row += 1

        block_end = current_row - 1
        # Auto-merge repeated Component / Rate Reference / Output cells,
        # same behaviour as the sample sheet (B5:B8, D5:D12, F5:F12 etc.)
        if block_end >= block_start:
            _merge_repeated_column(ws, "B", block_start, block_end)
            _merge_repeated_column(ws, "D", block_start, block_end)
            _merge_repeated_column(ws, "F", block_start, block_end)


def _write_rating_tables_sheet(ws: Worksheet, tables: List[Dict[str, Any]]) -> None:
    """
    Writes Sheet 2 exactly like the sample 'Rating Tables' tab: several
    lookup tables placed side-by-side, each with a merged title band,
    a colored header row, and bordered data cells. One blank column is
    left between tables (same visual gap as the sample).
    """
    TITLE_ROW = 2
    HEADER_ROW = 4
    DATA_START_ROW = 5

    current_col = 2  # start at column B, like the sample

    for table in tables:
        table_name = table.get("table_name", "")
        columns = table.get("columns", [])
        rows = table.get("rows", [])
        n_cols = max(1, len(columns))

        start_col_letter = get_column_letter(current_col)
        end_col_letter = get_column_letter(current_col + n_cols - 1)

        # Table title band, merged across the table's width
        if n_cols > 1:
            ws.merge_cells(f"{start_col_letter}{TITLE_ROW}:{end_col_letter}{TITLE_ROW}")
        title_cell = ws[f"{start_col_letter}{TITLE_ROW}"]
        title_cell.value = table_name
        title_cell.font = COLUMN_HEADER_FONT
        title_cell.fill = TABLE_TITLE_FILL
        title_cell.alignment = WRAP_CENTER
        title_cell.border = THIN_BORDER

        # Header row
        for offset, col_name in enumerate(columns):
            col_letter = get_column_letter(current_col + offset)
            cell = ws[f"{col_letter}{HEADER_ROW}"]
            cell.value = col_name
            cell.font = COLUMN_HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = WRAP_CENTER
            cell.border = THIN_BORDER
            ws.column_dimensions[col_letter].width = max(14, len(str(col_name)) + 4)

        # Data rows
        for r_offset, row_values in enumerate(rows):
            excel_row = DATA_START_ROW + r_offset
            for c_offset in range(n_cols):
                col_letter = get_column_letter(current_col + c_offset)
                cell = ws[f"{col_letter}{excel_row}"]
                cell.value = row_values[c_offset] if c_offset < len(row_values) else ""
                cell.font = BODY_FONT
                cell.alignment = WRAP_LEFT
                cell.border = THIN_BORDER

        # Move to next table's starting column, leaving 1 blank gap column
        current_col += n_cols + 1


def build_rating_workbook(product_name: str, extracted_data: Dict[str, Any]) -> Workbook:
    """
    Main entry point. Takes the structured dict returned by
    product_rating_extractor.extract_structured_rating() and returns
    an openpyxl Workbook styled exactly like the sample file.

    extracted_data = {"steps": [...], "tables": [...]}   (see contract above)
    """
    wb = Workbook()

    steps_ws = wb.active
    steps_ws.title = "Rating Steps"
    _write_rating_steps_sheet(steps_ws, product_name, extracted_data.get("steps", []))

    tables_ws = wb.create_sheet("Rating Tables")
    _write_rating_tables_sheet(tables_ws, extracted_data.get("tables", []))

    return wb


def save_workbook(wb: Workbook, output_path: str) -> str:
    """Saves the workbook to disk and returns the path (for CLI / Flask use)."""
    wb.save(output_path)
    return output_path